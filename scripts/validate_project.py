#!/usr/bin/env python3
"""Validate the retail sales source workbook and packaged Tableau dashboard."""

from __future__ import annotations

import math
import sys
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE_FILE = ROOT / "data" / "retail_sales_source.xlsx"
TABLEAU_FILE = ROOT / "tableau" / "retail_sales_dashboard.twbx"
EXPECTED_SHEETS = ("Order-Sales Data", "Metrics", "Cost Data", "Waffle Chart")
ORDER_COLUMNS = (
    "Order ID",
    "Order Date",
    "Ship Date",
    "Segment",
    "Category",
    "Subcategory",
    "Sales",
    "Customer ID",
    "Customer Name",
    "Country",
    "Returns",
    "Status",
    "Repeat Customer",
    "On Time",
)
EXPECTED_DASHBOARDS = {"Overview", "Sales", "Orders"}
EXPECTED_TABLEAU_DEFINITION = "retail_sales_dashboard.twb"


class ValidationError(Exception):
    """Raised when a required project artifact or data check is invalid."""


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValidationError(message)


def validate_source() -> dict[str, object]:
    require(SOURCE_FILE.is_file(), f"Missing source workbook: {SOURCE_FILE.relative_to(ROOT)}")
    workbook = load_workbook(SOURCE_FILE, read_only=True, data_only=True)
    try:
        require(
            tuple(workbook.sheetnames) == EXPECTED_SHEETS,
            f"Expected sheets {EXPECTED_SHEETS}; found {tuple(workbook.sheetnames)}",
        )

        order_sheet = workbook["Order-Sales Data"]
        rows = order_sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        require(header_row is not None, "Order-Sales Data sheet is empty")
        headers = tuple(value for value in header_row if value is not None)
        require(headers == ORDER_COLUMNS, f"Unexpected order columns: {headers}")
        require(
            all(value is None for value in header_row[len(ORDER_COLUMNS) :]),
            "Order-Sales Data has an unexpected named trailing column",
        )

        record_count = 0
        sales_total = 0.0
        order_ids: set[str] = set()
        customer_ids: set[str] = set()
        categories: set[str] = set()
        countries: set[str] = set()
        statuses: Counter[str] = Counter()
        segments: set[str] = set()
        first_order: datetime | None = None
        last_order: datetime | None = None

        for row_number, row in enumerate(rows, start=2):
            values = row[: len(ORDER_COLUMNS)]
            if not any(value is not None for value in values):
                continue
            require(
                all(value is None for value in row[len(ORDER_COLUMNS) :]),
                f"Order row {row_number} has values in the unnamed trailing column",
            )
            record = dict(zip(ORDER_COLUMNS, values))
            missing = [name for name, value in record.items() if value is None]
            require(not missing, f"Order row {row_number} has missing values: {missing}")

            order_date = record["Order Date"]
            ship_date = record["Ship Date"]
            require(isinstance(order_date, datetime), f"Order row {row_number} has an invalid Order Date")
            require(isinstance(ship_date, datetime), f"Order row {row_number} has an invalid Ship Date")
            require(ship_date >= order_date, f"Order row {row_number} ships before it was ordered")

            sales = record["Sales"]
            require(
                isinstance(sales, (int, float)) and math.isfinite(sales) and sales >= 0,
                f"Order row {row_number} has invalid Sales: {sales!r}",
            )
            for flag in ("Returns", "Repeat Customer", "On Time"):
                require(record[flag] in (0, 1, 0.0, 1.0), f"Order row {row_number} has invalid {flag}")

            record_count += 1
            sales_total += float(sales)
            order_ids.add(str(record["Order ID"]))
            customer_ids.add(str(record["Customer ID"]))
            categories.add(str(record["Category"]))
            countries.add(str(record["Country"]))
            segments.add(str(record["Segment"]))
            statuses[str(record["Status"])] += 1
            first_order = order_date if first_order is None else min(first_order, order_date)
            last_order = order_date if last_order is None else max(last_order, order_date)

        require(record_count == 21_900, f"Expected 21,900 order records; found {record_count}")
        require(len(categories) == 3, f"Expected 3 product categories; found {sorted(categories)}")
        require(len(countries) == 10, f"Expected 10 countries; found {len(countries)}")
        require(segments == {"Consumer", "Corporate", "Home"}, f"Unexpected segments: {sorted(segments)}")

        metrics = list(workbook["Metrics"].iter_rows(values_only=True))
        require(len(metrics) == 13, f"Expected 12 metric rows plus header; found {len(metrics) - 1}")
        require(
            {row[0] for row in metrics[1:]} == {"GPM", "OPM", "ROA", "ROE"},
            "Metrics sheet must contain GPM, OPM, ROA, and ROE",
        )

        costs = list(workbook["Cost Data"].iter_rows(values_only=True))
        require(len(costs) == 61, f"Expected 60 cost rows plus header; found {len(costs) - 1}")
        require(
            {row[2] for row in costs[1:]} == {"Shipping", "Product", "Operational", "Return", "Marketing"},
            "Cost Data sheet has an unexpected cost type set",
        )

        waffle = list(workbook["Waffle Chart"].iter_rows(values_only=True))
        require(len(waffle) == 101, f"Expected 100 waffle coordinates plus header; found {len(waffle) - 1}")
        coordinates = {(row[0], row[1]) for row in waffle[1:]}
        require(len(coordinates) == 100, "Waffle Chart must contain 100 unique coordinates")

        return {
            "records": record_count,
            "distinct_order_ids": len(order_ids),
            "customers": len(customer_ids),
            "sales_total": sales_total,
            "first_order": first_order,
            "last_order": last_order,
            "statuses": statuses,
        }
    finally:
        workbook.close()


def validate_tableau_package() -> None:
    require(TABLEAU_FILE.is_file(), f"Missing Tableau package: {TABLEAU_FILE.relative_to(ROOT)}")
    try:
        with zipfile.ZipFile(TABLEAU_FILE) as package:
            names = set(package.namelist())
            workbook_names = [name for name in names if name.lower().endswith(".twb")]
            require(len(workbook_names) == 1, f"Expected one .twb definition; found {workbook_names}")
            require(
                workbook_names[0] == EXPECTED_TABLEAU_DEFINITION,
                f"Expected internal workbook {EXPECTED_TABLEAU_DEFINITION}; found {workbook_names[0]}",
            )
            require(
                any(name.lower().endswith(".hyper") for name in names),
                "Tableau package does not contain a Hyper extract",
            )

            root = ElementTree.fromstring(package.read(workbook_names[0]))
            dashboards = {item.attrib.get("name") for item in root.findall(".//dashboard")}
            require(
                EXPECTED_DASHBOARDS <= dashboards,
                f"Missing expected dashboards: {sorted(EXPECTED_DASHBOARDS - dashboards)}",
            )

            extract_paths = {
                item.attrib["dbname"]
                for item in root.findall(".//connection")
                if item.attrib.get("class") == "hyper" and item.attrib.get("dbname")
            }
            missing = sorted(path for path in extract_paths if path not in names)
            require(not missing, f"Packaged workbook references missing extracts: {missing}")
    except (zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise ValidationError(f"Tableau package is not readable: {exc}") from exc


def main() -> int:
    try:
        summary = validate_source()
        validate_tableau_package()
    except (ValidationError, OSError, KeyError, StopIteration) as exc:
        print(f"Validation failed: {exc}", file=sys.stderr)
        return 1

    statuses = summary["statuses"]
    print("Project validation passed.")
    print(f"Order records: {summary['records']:,} ({summary['distinct_order_ids']:,} distinct order IDs)")
    print(f"Customers: {summary['customers']:,}")
    print(f"Order date range: {summary['first_order']:%Y-%m-%d} to {summary['last_order']:%Y-%m-%d}")
    print(f"Sales total: ${summary['sales_total']:,.2f}")
    print("Statuses: " + ", ".join(f"{name} {count:,}" for name, count in sorted(statuses.items())))
    print("Tableau package: workbook XML, expected dashboards, and referenced extracts present")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
